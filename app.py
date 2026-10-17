"""
Dashboard de Estudos · PMPE & PCPE 2026/2027
Rodar (dentro da pasta):  streamlit run app.py

Fontes de dados (mesma pasta deste arquivo):
- "CRONOGRAMA - ASSUNTO.xlsx": progresso. Toda edição feita no app é gravada nela.
- "cronograma - MATÉRIAS.xlsx": grade semanal usada no Plano de hoje (só leitura).
- pasta "base_estudos": resumos e flashcards (.md, um arquivo por matéria) e suas anotações.
"""
import base64
import datetime as dt
import json
import re
import shutil
from pathlib import Path

import pandas as pd
import requests
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st
import streamlit.components.v1 as components
from openpyxl import load_workbook

# ───────────────────────── Configuração ─────────────────────────
BASE = Path(__file__).parent
ARQ = BASE / "CRONOGRAMA - ASSUNTO.xlsx"
ARQ_FASES = BASE / "cronograma - MATÉRIAS.xlsx"
PASTA_BACKUP = BASE / "backup"
PASTA_ESTUDOS = BASE / "base_estudos"
ARQ_ANOTACOES = PASTA_ESTUDOS / "anotacoes.json"
ARQ_BANCA = PASTA_ESTUDOS / "analise_banca.md"
PASTA_MATERIAIS = BASE / "materiais"  # PDFs baixados do curso: ficam só no computador (fora do GitHub)
PERFIL_BANCA = {
    "AOCP": "Múltipla escolha com 5 alternativas, e o erro não anula acerto. Cobra muito a letra da lei, "
            "conceitos diretos e \"de acordo com o texto\". Leia todas as alternativas.",
    "Cebraspe": "Certo/Errado, e cada erro anula um acerto. Itens longos com pegadinhas em palavras como "
                "\"sempre\", \"somente\" e \"exclusivamente\". Gosta de jurisprudência e de reescrita. Na dúvida, deixe em branco.",
}
ABA_REGISTRO = "Registro_Estudos"
ABA_CONFIG = "Config_App"
ABA_PLANO = "Plano_Diario"
ABA_DASH = "Dashboard_Estratégico"
LINHA_INICIAL = 5  # 1ª linha de tópicos (cabeçalho na linha 4)

VINHO, VINHO_CLARO, ROSE, PRETO, TEXTO, APAGADO = "#8E2436", "#C04A5E", "#E8C3CA", "#0A0708", "#EFE6E7", "#A88E93"

# Colunas da planilha (posição 1 = coluna A)
COLUNAS = {
    "Cód": 1, "Cargos": 2, "Tópico": 3, "Fase": 4,
    "Data Planejada": 5, "Data Estudada": 6,
    "R1 (24h)": 7, "R2 (7d)": 8, "R3 (15d)": 9, "R4 (30d)": 10,
    "Nº Questões": 11, "% Acerto": 12, "Status": 13, "Observações": 14,
}
COLS_DATA = ["Data Planejada", "Data Estudada"]
REVISOES = [("R1 (24h)", 1), ("R2 (7d)", 7), ("R3 (15d)", 15), ("R4 (30d)", 30)]
STATUS = ["Não Iniciado", "Em Andamento", "Concluído", "Revisado"]
OPC_REV = ["Pendente", "Feito", "N/A"]
FEITO = {"Concluído", "Revisado"}

# Cada matéria (aba da planilha) tem nome curto, cor e marcador
MATERIAS = {
    "Português_e_Redação": ("Português & Redação", "#E0B84A", "🟡"),
    "Raciocínio_e_Estatística": ("Raciocínio & Estatística", "#4FA3E0", "🔵"),
    "Informática_e_Arquivologia": ("Informática & Arquivologia", "#5DBB63", "🟢"),
    "Direito_Constitucional_e_DH": ("Dir. Constitucional & DH", "#E05A5A", "🔴"),
    "Direito_Administrativo": ("Dir. Administrativo", "#A66CD1", "🟣"),
    "Direito_Penal_e_Processual": ("Dir. Penal & Processual", "#F08A3C", "🟠"),
    "Legislação_e_História_PE": ("Legislação & História PE", "#A47551", "🟤"),
    "Matérias_Específicas_Outras": ("Específicas (Oficial)", "#D9D9D9", "⚪"),
}

CARGOS = ["Soldado", "Oficial", "Agente", "Escrivão", "Delegado"]
# Cada concurso junta os cargos da mesma prova: o tópico entra se valer para qualquer um deles.
CONCURSOS = {
    "PMPE": {"nome": "PMPE · Oficial & Soldado", "cargos": ["Oficial", "Soldado"], "icone": "🎖️",
             "prova": dt.date(2027, 2, 21), "banca": "Instituto AOCP"},
    "PCPE": {"nome": "PCPE · Agente & Escrivão", "cargos": ["Agente", "Escrivão"], "icone": "🔎",
             "prova": dt.date(2027, 3, 7), "banca": "Cebraspe"},
}

st.set_page_config(page_title="Dashboard de Estudos", page_icon="📚", layout="wide")

# ───────────────────────── Estilo preto & vinho ─────────────────────────
st.markdown(f"""
<style>
@import url('https://fonts.googleapis.com/css2?family=Playfair+Display:wght@600;700&family=Inter:wght@400;500;600&display=swap');
html, body, [class*="css"], .stMarkdown, p, label {{ font-family: 'Inter', sans-serif; }}
.stApp {{ background: radial-gradient(circle at 15% 0%, #2a0d14 0%, {PRETO} 50%); }}
h1, h2, h3 {{ font-family: 'Playfair Display', serif !important; color: {ROSE} !important; letter-spacing: .3px; }}
.block-container {{ padding-top: 3.2rem; max-width: 1250px; }}
.titulo {{ font-family: 'Playfair Display', serif; font-size: 2.2rem; font-weight: 700; color: {ROSE}; margin-bottom: 0; }}
.titulo span {{ color: {VINHO_CLARO}; }}
.sub {{ color: {APAGADO}; font-size: .95rem; margin-top: -4px; }}
.kpi {{ background: linear-gradient(160deg, #1d0b10 0%, #0f0b0c 100%); border: 1px solid #4a1a24;
  border-radius: 16px; padding: 18px 20px; height: 100%; box-shadow: 0 8px 24px rgba(0,0,0,.45); }}
.kpi .rot {{ color: {APAGADO}; font-size: .76rem; text-transform: uppercase; letter-spacing: 1.4px; }}
.kpi .val {{ color: {ROSE}; font-family: 'Playfair Display', serif; font-size: 2rem; font-weight: 700; line-height: 1.15; }}
.kpi .det {{ color: {TEXTO}; font-size: .85rem; opacity: .75; }}
.barra {{ height: 6px; background: #2a1218; border-radius: 6px; margin-top: 10px; overflow: hidden; }}
.barra > div {{ height: 100%; background: linear-gradient(90deg, #5c1422, {VINHO}, {VINHO_CLARO}); }}
.chip {{ display: inline-block; border: 1px solid #5a2030; border-radius: 999px; padding: 4px 12px; margin: 4px 6px 0 0;
  color: {TEXTO}; font-size: .82rem; background: #140a0d; }}
.chip b {{ color: {VINHO_CLARO}; }}
.resumo {{ background: #120c0e; border: 1px solid #3a1820; border-radius: 14px; padding: 18px 22px; color: {TEXTO}; line-height: 1.6; }}
.flash {{ background: linear-gradient(160deg, #2a0f17, #120a0c); border: 1px solid {VINHO}; border-radius: 18px;
  padding: 34px 28px; text-align: center; font-size: 1.15rem; color: {TEXTO}; min-height: 150px; }}
.flash small {{ display:block; color: {APAGADO}; font-size: .75rem; letter-spacing: 1.4px; text-transform: uppercase; margin-bottom: 12px; }}
.stButton > button, .stFormSubmitButton > button {{ border: 1px solid {VINHO}; color: {ROSE}; background: transparent; border-radius: 10px; }}
.stButton > button:hover, .stFormSubmitButton > button:hover {{ background: {VINHO}; color: #fff; border-color: {VINHO}; }}
.stButton > button[kind="primary"] {{ background: {VINHO}; color: #fff; border-color: {VINHO}; }}
[data-testid="stSidebar"] {{ background: #0d0809; border-right: 1px solid #3a1820; }}
[data-testid="stSidebar"] .stButton > button {{ border: none; justify-content: flex-start; text-align: left; color: {TEXTO}; }}
[data-testid="stSidebar"] .stButton > button[kind="primary"] {{ background: #3a121c; color: {ROSE}; border-left: 3px solid {VINHO_CLARO}; }}
[data-testid="stSidebar"] .stButton > button:hover {{ background: #2a0f17; color: {ROSE}; }}
[data-testid="stExpander"] {{ border: 1px solid #3a1820; border-radius: 12px; }}
[data-testid="stCheckbox"] label p {{ font-size: .95rem; }}
hr {{ border-color: #3a1820 !important; }}
</style>
""", unsafe_allow_html=True)


# ───────────────────────── Funções auxiliares ─────────────────────────
def vazio(v):
    if v is None:
        return True
    if isinstance(v, str):
        return v.strip() == ""
    try:
        return bool(pd.isna(v))
    except (TypeError, ValueError):
        return False


def para_data(v):
    if vazio(v):
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)):  # número serial do Excel
        return dt.date(1899, 12, 30) + dt.timedelta(days=int(v))
    txt = str(v).strip()[:10]
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return dt.datetime.strptime(txt, fmt).date()
        except ValueError:
            pass
    return None


def para_num(v):
    if vazio(v):
        return 0.0
    try:
        return float(str(v).replace("%", "").replace(",", "."))
    except ValueError:
        return 0.0


def acerto_fracao(v):
    """Aceita 0,8 / 80 / '80%' e devolve 0.8."""
    n = para_num(v)
    if isinstance(v, str) and "%" in v:
        n = n / 100
    elif n > 1:
        n = n / 100
    return max(0.0, min(1.0, n))


def cargos_do_topico(txt):
    t = str(txt or "")
    if "todos" in t.lower():
        return set(CARGOS)
    achados = set()
    for tok in re.split(r"[/()\s,]+", t):
        tok = tok.strip().lower()
        if tok.startswith("sold"):
            achados.add("Soldado")
        elif tok.startswith("ofic"):
            achados.add("Oficial")
        elif tok.startswith("agen"):
            achados.add("Agente")
        elif tok.startswith("escr"):
            achados.add("Escrivão")
        elif tok.startswith("del"):
            achados.add("Delegado")
    if not achados:
        if "PMPE" in t:
            achados |= {"Soldado", "Oficial"}
        if "PCPE" in t:
            achados |= {"Agente", "Escrivão", "Delegado"}
    return achados


def info_materia(aba):
    if aba in MATERIAS:
        return MATERIAS[aba]
    extras = ["#7FD1C7", "#C9A0DC", "#F5B7B1", "#AED6F1"]
    return (aba.replace("_", " "), extras[sum(map(ord, aba)) % len(extras)], "⚫")


def abas_de_materia(wb):
    """Toda aba com 'Cód' em A4 é tratada como matéria (novas abas entram sozinhas)."""
    return [ws.title for ws in wb.worksheets if str(ws.cell(4, 1).value or "").strip() == "Cód"]


# ───────────────────────── Leitura ─────────────────────────
@st.cache_data(show_spinner=False)
def carregar(_mtime):
    wb = load_workbook(ARQ, data_only=True)
    linhas = []
    for aba in abas_de_materia(wb):
        ws = wb[aba]
        nome, cor, emoji = info_materia(aba)
        for r in range(LINHA_INICIAL, ws.max_row + 1):
            cod = ws.cell(r, 1).value
            if vazio(cod):
                continue
            reg = {"id": f"{aba}|{r}", "aba": aba, "linha": r,
                   "Matéria": f"{emoji} {nome}", "MatériaNome": nome, "Cor": cor}
            for col, idx in COLUNAS.items():
                reg[col] = ws.cell(r, idx).value
            reg["Cód"] = str(reg["Cód"]).strip()
            reg["Data Planejada"] = para_data(reg["Data Planejada"])
            reg["Data Estudada"] = para_data(reg["Data Estudada"])
            reg["Nº Questões"] = int(para_num(reg["Nº Questões"]))
            reg["% Acerto"] = round(acerto_fracao(reg["% Acerto"]) * 100, 1)
            reg["Status"] = reg["Status"] if reg["Status"] in STATUS else "Não Iniciado"
            for c, _ in REVISOES:
                reg[c] = reg[c] if reg[c] in OPC_REV else "Pendente"
            reg["Observações"] = "" if vazio(reg["Observações"]) else str(reg["Observações"])
            cs = cargos_do_topico(reg["Cargos"])
            for c in CARGOS:
                reg[c] = c in cs
            for k, conc in CONCURSOS.items():
                reg[k] = any(c in cs for c in conc["cargos"])
            linhas.append(reg)
    topicos = pd.DataFrame(linhas).set_index("id")

    registros = pd.DataFrame(columns=["Data", "Matéria", "Cód", "Minutos", "Tipo"])
    if ABA_REGISTRO in wb.sheetnames:
        ws = wb[ABA_REGISTRO]
        dados = [r for r in ws.iter_rows(min_row=2, values_only=True) if r and not vazio(r[0])]
        if dados:
            registros = pd.DataFrame([list(r[:5]) + [None] * (5 - len(r[:5])) for r in dados],
                                     columns=registros.columns)
            registros["Data"] = registros["Data"].map(para_data)
            registros["Minutos"] = registros["Minutos"].map(para_num)
            registros = registros.dropna(subset=["Data"])

    plano_feito = {}
    if ABA_PLANO in wb.sheetnames:
        for d, item, feito in wb[ABA_PLANO].iter_rows(min_row=2, max_col=3, values_only=True):
            if para_data(d) and item:
                plano_feito[(para_data(d), str(item))] = bool(feito)

    config = {"meta": 27.0, "inicio": None}
    if ABA_CONFIG in wb.sheetnames:
        for chave, valor in wb[ABA_CONFIG].iter_rows(min_row=1, max_col=2, values_only=True):
            if chave == "Meta semanal (h)":
                config["meta"] = para_num(valor) or 27.0
            elif chave == "Início do plano":
                config["inicio"] = para_data(valor)
    if config["inicio"] is None:
        datas = list(registros["Data"]) + [d for d in topicos["Data Estudada"] if d]
        config["inicio"] = min(datas) if datas else dt.date.today()
    return topicos, registros, plano_feito, config


@st.cache_data(show_spinner=False)
def plano_semanal(_mtime):
    """Lê a grade semanal da Fase 1 em 'cronograma - MATÉRIAS.xlsx'."""
    if not ARQ_FASES.exists():
        return {}
    try:
        wb = load_workbook(ARQ_FASES, data_only=True)
    except Exception:
        return {}
    grade = {}
    for ws in wb.worksheets:
        cab = None
        for row in ws.iter_rows(values_only=True):
            if row and str(row[0] or "").strip() == "Dia da Semana":
                cab = [str(c or "") for c in row]
                continue
            if cab and row and row[0] and str(row[0]).endswith(("-feira", "Sábado", "Domingo")):
                grade[str(row[0]).strip()] = dict(zip(cab, row))
    return grade


@st.cache_data(show_spinner=False)
def carregar_base(_assinatura):
    """Lê base_estudos/*.md. Formato: '## CÓD', '### Resumo', '### Flashcards' com linhas 'P:' e 'R:'."""
    base = {}
    for arq in sorted(PASTA_ESTUDOS.glob("*.md")):
        if arq.name == ARQ_BANCA.name:
            continue
        texto = arq.read_text(encoding="utf-8")
        for bloco in re.split(r"^## ", texto, flags=re.M)[1:]:
            cod, _, corpo = bloco.partition("\n")
            resumo, _, cartas = corpo.partition("### Flashcards")
            resumo = re.sub(r"^### Resumo\s*", "", resumo.strip()).strip()
            flash = re.findall(r"^\s*-?\s*P:\s*(.+?)\s*\n\s*R:\s*(.+?)\s*$", cartas, flags=re.M)
            base[cod.strip()] = {"resumo": resumo, "flashcards": flash}
    return base


@st.cache_data(show_spinner=False)
def carregar_banca(_mtime):
    """Lê base_estudos/analise_banca.md: '## CÓD' + linhas 'Palavras-chave:', 'AOCP:', 'Cebraspe:'."""
    if not ARQ_BANCA.exists():
        return {}
    dados = {}
    for bloco in re.split(r"^## ", ARQ_BANCA.read_text(encoding="utf-8"), flags=re.M)[1:]:
        cod, _, corpo = bloco.partition("\n")
        item = {"palavras": [], "AOCP": "", "Cebraspe": ""}
        for linha in corpo.splitlines():
            chave, sep, valor = linha.partition(":")
            if not sep:
                continue
            if chave.strip() == "Palavras-chave":
                item["palavras"] = [p.strip() for p in valor.split(";") if p.strip()]
            elif chave.strip() in ("AOCP", "Cebraspe"):
                item[chave.strip()] = valor.strip()
        dados[cod.strip()] = item
    return dados


def pontos_do_resumo(resumo):
    """Cada tópico (linha com '-') do resumo vira um ponto para anotar."""
    return [l.strip()[2:].strip() for l in resumo.splitlines() if l.strip().startswith("- ")]


def ler_anotacoes():
    try:
        return json.loads(ARQ_ANOTACOES.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return {}


# ───────────────────────── Sincronização com o GitHub (versão online) ─────────────────────────
# Online (Streamlit Cloud) os arquivos do servidor são apagados quando o app reinicia. Por isso, se houver
# um token em Secrets ([github] token = "..."), a planilha e as anotações são lidas e gravadas no branch
# "dados" do repositório. Sem token (no computador), nada disso acontece.
def config_github():
    try:
        cfg = st.secrets["github"]
        return {"token": cfg["token"], "repo": cfg.get("repo", "cristielesilva/dashboard-estudos"),
                "branch": cfg.get("branch", "dados")}
    except Exception:
        return None


def _api(cfg, metodo, caminho, **kw):
    cab = {"Authorization": f"Bearer {cfg['token']}", "Accept": "application/vnd.github+json"}
    cab.update(kw.pop("headers", {}))
    url = f"https://api.github.com/repos/{cfg['repo']}" + (f"/{caminho}" if caminho else "")
    return requests.request(metodo, url,
                            headers=cab, timeout=20, **kw)


class ErroGitHub(Exception):
    pass


def _garantir_branch(cfg):
    repo = _api(cfg, "GET", "")
    if repo.status_code in (401, 403, 404):
        raise ErroGitHub(
            f"O GitHub recusou a chave (código {repo.status_code}). Confira em Settings > Secrets se o token "
            f"foi colado inteiro, entre aspas, e se ele tem acesso ao repositório {cfg['repo']} "
            "com a permissão Contents em Read and write.")
    if _api(cfg, "GET", f"branches/{cfg['branch']}").status_code == 200:
        return
    padrao = repo.json().get("default_branch", "main")
    ref = _api(cfg, "GET", f"git/ref/heads/{padrao}")
    criar = None
    if ref.status_code == 200:
        criar = _api(cfg, "POST", "git/refs",
                     json={"ref": f"refs/heads/{cfg['branch']}", "sha": ref.json()["object"]["sha"]})
    if criar is None or criar.status_code not in (201, 422):
        raise ErroGitHub(
            f"Não consegui criar o branch \"{cfg['branch']}\" no GitHub. Verifique se o token tem a "
            "permissão Contents em Read and write.")


@st.cache_resource(show_spinner="Carregando seus dados salvos...")
def baixar_do_github():
    """Roda uma vez quando o servidor liga: traz a versão mais recente dos dados."""
    cfg = config_github()
    if not cfg:
        return None
    try:
        _garantir_branch(cfg)
    except (ErroGitHub, requests.RequestException) as erro:
        return str(erro)
    for arq in (ARQ, ARQ_ANOTACOES):
        rel = arq.relative_to(BASE).as_posix()
        r = _api(cfg, "GET", f"contents/{rel}", params={"ref": cfg["branch"]},
                 headers={"Accept": "application/vnd.github.raw"})
        if r.status_code == 200:
            arq.parent.mkdir(exist_ok=True)
            arq.write_bytes(r.content)
    return None


def enviar_ao_github(arq):
    cfg = config_github()
    if not cfg:
        return
    rel = arq.relative_to(BASE).as_posix()
    try:
        atual = _api(cfg, "GET", f"contents/{rel}", params={"ref": cfg["branch"]})
        corpo = {"message": f"Atualiza {arq.name} pelo dashboard", "branch": cfg["branch"],
                 "content": base64.b64encode(arq.read_bytes()).decode()}
        if atual.status_code == 200:
            corpo["sha"] = atual.json()["sha"]
        r = _api(cfg, "PUT", f"contents/{rel}", json=corpo)
        r.raise_for_status()
    except Exception as erro:
        st.warning(f"Salvo no app, mas não consegui enviar ao GitHub ({erro}). Confira o token em Secrets.")


# ───────────────────────── Gravação ─────────────────────────
def fazer_backup():
    if st.session_state.get("backup_feito"):
        return
    PASTA_BACKUP.mkdir(exist_ok=True)
    carimbo = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    shutil.copy2(ARQ, PASTA_BACKUP / f"CRONOGRAMA - ASSUNTO_{carimbo}.xlsx")
    for velho in sorted(PASTA_BACKUP.glob("CRONOGRAMA - ASSUNTO_*.xlsx"))[:-10]:  # guarda os 10 últimos
        velho.unlink(missing_ok=True)
    st.session_state.backup_feito = True


def atualizar_dashboard(wb):
    """Atualiza a coluna '% Concluído' da aba Dashboard_Estratégico."""
    if ABA_DASH not in wb.sheetnames:
        return
    total = {c: 0 for c in CARGOS}
    feitos = {c: 0 for c in CARGOS}
    for aba in abas_de_materia(wb):
        ws = wb[aba]
        for r in range(LINHA_INICIAL, ws.max_row + 1):
            if vazio(ws.cell(r, 1).value):
                continue
            ok = ws.cell(r, COLUNAS["Status"]).value in FEITO
            for c in cargos_do_topico(ws.cell(r, COLUNAS["Cargos"]).value):
                total[c] += 1
                feitos[c] += ok
    ws = wb[ABA_DASH]
    for r in range(1, ws.max_row + 1):
        nome = str(ws.cell(r, 1).value or "")
        for c in CARGOS:
            if c.lower() in nome.lower() and total[c]:
                cel = ws.cell(r, 8)
                cel.value = feitos[c] / total[c]
                cel.number_format = "0.0%"


def salvar(alteracoes=(), novos_registros=(), config=None, plano=None):
    """alteracoes: lista de (id, coluna, valor). Grava tudo de uma vez na planilha original."""
    try:
        fazer_backup()
        wb = load_workbook(ARQ)
        for ident, col, valor in alteracoes:
            aba, linha = ident.rsplit("|", 1)
            cel = wb[aba].cell(int(linha), COLUNAS[col])
            if col in COLS_DATA:
                cel.value = valor
                if valor:
                    cel.number_format = "DD/MM/YYYY"
            elif col == "% Acerto":
                cel.value = round((valor or 0) / 100, 4)
                cel.number_format = "0.0%"
            elif col == "Nº Questões":
                cel.value = int(valor or 0)
            else:
                cel.value = valor
        if novos_registros:
            if ABA_REGISTRO not in wb.sheetnames:
                wb.create_sheet(ABA_REGISTRO).append(["Data", "Matéria", "Cód", "Minutos", "Tipo"])
            ws = wb[ABA_REGISTRO]
            for reg in novos_registros:
                ws.append(reg)
                ws.cell(ws.max_row, 1).number_format = "DD/MM/YYYY"
        if plano:
            data, item, feito = plano
            if ABA_PLANO not in wb.sheetnames:
                wb.create_sheet(ABA_PLANO).append(["Data", "Item do plano", "Feito"])
            ws = wb[ABA_PLANO]
            for r in range(2, ws.max_row + 1):
                if para_data(ws.cell(r, 1).value) == data and ws.cell(r, 2).value == item:
                    ws.cell(r, 3).value = feito
                    break
            else:
                ws.append([data, item, feito])
                ws.cell(ws.max_row, 1).number_format = "DD/MM/YYYY"
        if config:
            if ABA_CONFIG in wb.sheetnames:
                del wb[ABA_CONFIG]
            ws = wb.create_sheet(ABA_CONFIG)
            ws.append(["Meta semanal (h)", config["meta"]])
            ws.append(["Início do plano", config["inicio"]])
            ws.cell(2, 2).number_format = "DD/MM/YYYY"
        atualizar_dashboard(wb)
        wb.save(ARQ)
        enviar_ao_github(ARQ)
    except PermissionError:
        st.error("⚠️ Não consegui salvar: a planilha está aberta no Excel. Feche o arquivo e repita a alteração.")
        return False
    st.cache_data.clear()
    st.session_state.versao = st.session_state.get("versao", 0) + 1
    st.session_state.ultimo_salvo = dt.datetime.now().strftime("%H:%M:%S")
    return True


def salvar_anotacao(cod, texto=None, links=None, pdfs=None):
    """Grava só os campos informados, mantendo os demais."""
    PASTA_ESTUDOS.mkdir(exist_ok=True)
    todas = ler_anotacoes()
    nota = todas.get(cod, {})
    for campo, valor in (("texto", texto), ("links", links), ("pdfs", pdfs)):
        if valor is not None:
            nota[campo] = valor
    nota["atualizado"] = dt.date.today().isoformat()
    todas[cod] = nota
    ARQ_ANOTACOES.write_text(json.dumps(todas, ensure_ascii=False, indent=2), encoding="utf-8")
    enviar_ao_github(ARQ_ANOTACOES)


# ───────────────────────── Componentes visuais ─────────────────────────
def kpi(rotulo, valor, detalhe="", progresso=None):
    barra = ""
    if progresso is not None:
        p = max(0, min(100, progresso))
        barra = f'<div class="barra"><div style="width:{p:.0f}%"></div></div>'
    st.markdown(f'<div class="kpi"><div class="rot">{rotulo}</div><div class="val">{valor}</div>'
                f'<div class="det">{detalhe}</div>{barra}</div>', unsafe_allow_html=True)


def estilo(fig, altura=330, titulo=None):
    fig.update_layout(template="plotly_dark", paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)",
                      font=dict(color=TEXTO, family="Inter"), height=altura, showlegend=False,
                      margin=dict(l=10, r=20, t=50 if titulo else 10, b=10),
                      title=dict(text=titulo, font=dict(color=ROSE, size=16, family="Playfair Display")))
    fig.update_xaxes(gridcolor="#2a1218", zerolinecolor="#2a1218")
    fig.update_yaxes(gridcolor="#2a1218", zerolinecolor="#2a1218")
    return fig


def cores(df):
    return dict(zip(df["Matéria"], df["Cor"]))


def taxa_acerto(df):
    q = df["Nº Questões"].sum()
    return (df["Nº Questões"] * df["% Acerto"]).sum() / q if q else 0.0


def fmt_p(p):
    return f"{p:.1f}%".replace(".", ",")


def fmt_h(h):
    return f"{h:,.1f}".replace(",", "X").replace(".", ",").replace("X", ".") + " h"


def avisar(msg, icone):
    """Guarda o aviso para mostrar depois do recarregamento da página."""
    st.session_state.aviso = (msg, icone)


def ir(pagina, **extra):
    st.session_state.pagina = pagina
    st.session_state.update(extra)


def resumo_materias(df):
    g = df.groupby(["Matéria", "Cor"], sort=False).agg(
        total=("Status", "size"), feitos=("Status", lambda s: s.isin(FEITO).sum()),
        questoes=("Nº Questões", "sum")).reset_index()
    g["% Concluído"] = (g["feitos"] / g["total"] * 100).round(1)
    g["% Acerto"] = [round(taxa_acerto(df[df["Matéria"] == m]), 1) for m in g["Matéria"]]
    return g


def revisoes_pendentes(df, ate):
    linhas = []
    for ident, r in df.iterrows():
        base = r["Data Estudada"]
        if not base:
            continue
        for col, dias in REVISOES:
            if r[col] == "Pendente":
                venc = base + dt.timedelta(days=dias)
                if venc <= ate:
                    linhas.append({"id": ident, "Coluna": col, "Matéria": r["Matéria"], "Cód": r["Cód"],
                                   "Tópico": r["Tópico"], "Revisão": col, "Vence em": venc,
                                   "Atraso (dias)": (dt.date.today() - venc).days})
    return pd.DataFrame(linhas)


# ───────────────────────── Tabela editável (salva automaticamente) ─────────────────────────
ORDEM = ["Status", "Matéria", "Cód", "Tópico", "Data Estudada", "Nº Questões", "% Acerto",
         "R1 (24h)", "R2 (7d)", "R3 (15d)", "R4 (30d)", "Data Planejada", "Cargos", "Observações"]
EDITAVEIS = ["Status", "Data Estudada", "Nº Questões", "% Acerto", "R1 (24h)", "R2 (7d)",
             "R3 (15d)", "R4 (30d)", "Data Planejada", "Observações"]


def normalizar(col, v):
    if col in COLS_DATA:
        return para_data(v)
    if col == "Nº Questões":
        return int(para_num(v))
    if col == "% Acerto":
        return round(max(0.0, min(100.0, para_num(v))), 1)
    return "" if vazio(v) else str(v)


def tabela_editavel(df, chave):
    vis = df[ORDEM].copy()
    for c in COLS_DATA:
        vis[c] = pd.to_datetime(vis[c])
    cfg = {
        "Status": st.column_config.SelectboxColumn("Status", options=STATUS, required=True, width="small"),
        "Matéria": st.column_config.TextColumn("Matéria", width="medium"),
        "Cód": st.column_config.TextColumn("Cód", width="small"),
        "Tópico": st.column_config.TextColumn("Tópico", width="large"),
        "Data Estudada": st.column_config.DateColumn("Estudado em", format="DD/MM/YYYY"),
        "Data Planejada": st.column_config.DateColumn("Planejado p/", format="DD/MM/YYYY"),
        "Nº Questões": st.column_config.NumberColumn("Questões", min_value=0, step=1),
        "% Acerto": st.column_config.NumberColumn("% Acerto", min_value=0, max_value=100, step=1, format="%.0f%%"),
        "Cargos": st.column_config.TextColumn("Cargos", width="small"),
        "Observações": st.column_config.TextColumn("Observações", width="medium"),
        **{c: st.column_config.SelectboxColumn(c, options=OPC_REV, width="small") for c, _ in REVISOES},
    }
    editado = st.data_editor(
        vis, key=f"{chave}_v{st.session_state.get('versao', 0)}", column_config=cfg, hide_index=True,
        width="stretch", height=min(38 * len(vis) + 40, 620),
        disabled=[c for c in ORDEM if c not in EDITAVEIS])

    alteracoes = []
    for ident in vis.index:
        for col in EDITAVEIS:
            antes, depois = normalizar(col, df.at[ident, col]), normalizar(col, editado.at[ident, col])
            if antes != depois:
                alteracoes.append((ident, col, depois))
                # Concluiu e não marcou a data? Preenche com hoje para as revisões começarem a contar.
                if col == "Status" and depois in FEITO and not df.at[ident, "Data Estudada"] \
                        and not normalizar("Data Estudada", editado.at[ident, "Data Estudada"]):
                    alteracoes.append((ident, "Data Estudada", dt.date.today()))
    if alteracoes and salvar(alteracoes):
        avisar(f"Salvo na planilha ✓ ({len(alteracoes)} alteração(ões))", "💾")
        st.rerun()


# ───────────────────────── Pomodoro (HTML/JS) ─────────────────────────
POMODORO_HTML = f"""
<div id="p" style="font-family:Inter,sans-serif;color:{TEXTO};text-align:center;padding:4px">
  <div style="display:flex;gap:8px;justify-content:center;margin-bottom:8px">
    <button class="m" data-m="25">Foco 25</button><button class="m" data-m="5">Pausa 5</button><button class="m" data-m="15">Pausa 15</button>
  </div>
  <svg width="200" height="200" viewBox="0 0 230 230">
    <circle cx="115" cy="115" r="100" stroke="#2a1218" stroke-width="12" fill="none"/>
    <circle id="arc" cx="115" cy="115" r="100" stroke="{VINHO_CLARO}" stroke-width="12" fill="none" stroke-linecap="round"
      transform="rotate(-90 115 115)" stroke-dasharray="628.3" stroke-dashoffset="0"/>
    <text id="t" x="115" y="125" text-anchor="middle" font-size="46" fill="{ROSE}" font-family="Playfair Display,serif">25:00</text>
    <text id="lbl" x="115" y="155" text-anchor="middle" font-size="13" fill="{APAGADO}">FOCO</text>
  </svg>
  <div style="display:flex;gap:10px;justify-content:center;margin-top:6px">
    <button id="go" class="b">▶ Iniciar</button><button id="rs" class="b">↺ Zerar</button>
  </div>
  <div id="msg" style="margin-top:8px;font-size:.85rem;color:{APAGADO}">Ciclos de foco nesta sessão: 0</div>
</div>
<style>
 button{{background:transparent;border:1px solid {VINHO};color:{ROSE};border-radius:10px;padding:6px 14px;cursor:pointer;font-size:.9rem}}
 button:hover,button.on{{background:{VINHO};color:#fff}}
</style>
<script>
const K='pomodoro_v1', C=628.3;
let s={{modo:25,rest:1500,rodando:false,fim:null,ciclos:0}}, ctx=null;
try{{const x=JSON.parse(localStorage.getItem(K)); if(x) s=x;}}catch(e){{}}
const $=id=>document.getElementById(id);
function grava(){{try{{localStorage.setItem(K,JSON.stringify(s));}}catch(e){{}}}}
function bip(){{try{{ctx=ctx||new AudioContext();[0,0.35,0.7].forEach(d=>{{const o=ctx.createOscillator(),g=ctx.createGain();
 o.frequency.value=880;o.connect(g);g.connect(ctx.destination);g.gain.setValueAtTime(.25,ctx.currentTime+d);
 g.gain.exponentialRampToValueAtTime(.001,ctx.currentTime+d+.3);o.start(ctx.currentTime+d);o.stop(ctx.currentTime+d+.3);}});}}catch(e){{}}}}
function modo(m){{s.modo=m;s.rest=m*60;s.rodando=false;s.fim=null;grava();desenha();}}
function desenha(){{
  const m=Math.floor(s.rest/60),sg=Math.floor(s.rest%60);
  $('t').textContent=String(m).padStart(2,'0')+':'+String(sg).padStart(2,'0');
  $('arc').style.strokeDashoffset=C*(1-s.rest/(s.modo*60));
  $('lbl').textContent=s.modo==25?'FOCO':'PAUSA';
  $('go').textContent=s.rodando?'❚❚ Pausar':'▶ Iniciar';
  document.querySelectorAll('.m').forEach(b=>b.classList.toggle('on',+b.dataset.m==s.modo));
  $('msg').textContent='Ciclos de foco nesta sessão: '+s.ciclos;
}}
setInterval(()=>{{ if(!s.rodando) return;
  s.rest=Math.max(0,(s.fim-Date.now())/1000);
  if(s.rest<=0){{ s.rodando=false; bip();
    if(s.modo==25){{s.ciclos++; const longa=s.ciclos%4==0; s.modo=longa?15:5; s.rest=s.modo*60;
      $('msg').textContent='Pomodoro concluído! Registre abaixo.';}}
    else {{ s.modo=25; s.rest=1500; }}
    grava(); desenha(); return; }}
  desenha(); }}, 250);
$('go').onclick=()=>{{ try{{ctx=ctx||new AudioContext();}}catch(e){{}}
  if(s.rodando){{s.rodando=false;}} else {{s.rodando=true;s.fim=Date.now()+s.rest*1000;}} grava();desenha(); }};
$('rs').onclick=()=>modo(s.modo);
document.querySelectorAll('.m').forEach(b=>b.onclick=()=>modo(+b.dataset.m));
desenha();
</script>
"""


def pomodoro():
    if hasattr(st, "iframe"):
        st.iframe(POMODORO_HTML, height=320)
    else:
        components.html(POMODORO_HTML, height=320)


def registrar_estudo(topicos, chave, cod_fixo=None):
    """Botões para lançar um pomodoro (25 min) ou minutos avulsos numa matéria/tópico."""
    if cod_fixo:
        linha = topicos[topicos["Cód"] == cod_fixo].iloc[0]
        nome_mat, cod, ident = linha["MatériaNome"], cod_fixo, topicos.index[topicos["Cód"] == cod_fixo][0]
    else:
        mats = list(dict.fromkeys(topicos["Matéria"]))
        mat = st.selectbox("Matéria estudada", mats, key=f"{chave}_mat")
        sub = topicos[topicos["Matéria"] == mat]
        opcoes = ["— sem tópico específico —"] + [f"{r['Cód']} · {str(r['Tópico'])[:60]}" for _, r in sub.iterrows()]
        top = st.selectbox("Tópico", opcoes, key=f"{chave}_top")
        nome_mat = sub["MatériaNome"].iloc[0]
        cod = "" if top.startswith("—") else top.split(" · ")[0]
        ident = sub.index[sub["Cód"] == cod][0] if cod else None
    b1, b2, b3 = st.columns([1.4, .8, 1])
    reg_pomo = b1.button("✅ Pomodoro 25 min", key=f"{chave}_pomo", width="stretch")
    minutos = b2.number_input("min", 5, 600, 60, 5, key=f"{chave}_min", label_visibility="collapsed")
    reg_man = b3.button(f"➕ {minutos} min", key=f"{chave}_man", width="stretch")
    if reg_pomo or reg_man:
        mins, tipo = (25, "Pomodoro") if reg_pomo else (minutos, "Manual")
        alt = []
        if ident and topicos.at[ident, "Status"] == "Não Iniciado":
            alt.append((ident, "Status", "Em Andamento"))
        if salvar(alt, [[dt.date.today(), nome_mat, cod, mins, tipo]]):
            avisar(f"{mins} min registrados em {nome_mat} ✓", "⏱️")
            st.rerun()


# ───────────────────────── Dados ─────────────────────────
erro_sync = baixar_do_github()
if erro_sync:
    st.warning(f"⚠️ Suas edições não estão sendo salvas online. {erro_sync}")
if not ARQ.exists():
    st.error(f"Não encontrei a planilha **{ARQ.name}** na pasta {BASE}. Coloque o app.py na mesma pasta dela.")
    st.stop()

todos, registros, plano_feito, config = carregar(ARQ.stat().st_mtime)
# Só entram os tópicos de PMPE (Oficial + Soldado) e PCPE (Agente + Escrivão); os exclusivos de Delegado ficam de fora.
topicos = todos[todos[list(CONCURSOS)].any(axis=1)]
grade = plano_semanal(ARQ_FASES.stat().st_mtime if ARQ_FASES.exists() else 0)
banca = carregar_banca(ARQ_BANCA.stat().st_mtime if ARQ_BANCA.exists() else 0)
base_estudos = carregar_base(tuple(f.stat().st_mtime for f in sorted(PASTA_ESTUDOS.glob("*.md"))))
hoje = dt.date.today()
st.session_state.setdefault("pagina", "inicio")
pagina = st.session_state.pagina

if "aviso" in st.session_state:
    _msg, _ic = st.session_state.pop("aviso")
    st.toast(_msg, icon=_ic)

# ───────────────────────── Menu lateral ─────────────────────────
with st.sidebar:
    st.markdown(f"<div class='titulo' style='font-size:1.5rem'>Meus <span>Estudos</span></div>", unsafe_allow_html=True)
    st.write("")
    st.button("🏠 Início", on_click=ir, args=("inicio",), width="stretch",
              type="primary" if pagina == "inicio" else "secondary")
    st.button("📚 Estudar", on_click=ir, args=("estudar",), width="stretch",
              type="primary" if pagina == "estudar" else "secondary")
    st.button("📝 Resumo do dia", on_click=ir, args=("resumo_dia",), width="stretch",
              type="primary" if pagina == "resumo_dia" else "secondary")
    n_rev = len(revisoes_pendentes(topicos, hoje))
    st.button(f"🔁 Revisões ({n_rev})", on_click=ir, args=("revisoes",), width="stretch",
              type="primary" if pagina == "revisoes" else "secondary")
    st.write("")
    for chave, conc in CONCURSOS.items():
        df_c = topicos[topicos[chave]]
        pct = df_c["Status"].isin(FEITO).mean() * 100 if len(df_c) else 0
        ativo = pagina == f"conc:{chave}"
        with st.expander(f"{conc['icone']} {chave} · {pct:.0f}%", expanded=ativo):
            mat_ativa = st.session_state.get("materia_sel") if ativo else None
            st.button("📋 Todas as matérias", key=f"nav_{chave}_todas", width="stretch",
                      on_click=ir, args=(f"conc:{chave}",), kwargs={"materia_sel": None},
                      type="primary" if ativo and not mat_ativa else "secondary")
            for mat, g in df_c.groupby("Matéria", sort=False):
                p = g["Status"].isin(FEITO).mean() * 100
                st.button(f"{mat} · {p:.0f}%", key=f"nav_{chave}_{mat}", width="stretch",
                          on_click=ir, args=(f"conc:{chave}",), kwargs={"materia_sel": mat},
                          type="primary" if ativo and mat_ativa == mat else "secondary")
    st.write("")
    with st.expander("⚙️ Metas e arquivo"):
        meta = st.number_input("Meta semanal de horas", 1.0, 80.0, float(config["meta"]), 0.5, format="%.1f",
                               help="A planilha sugere 24h–30h/semana na Fase 1.")
        inicio = st.date_input("Início do plano", config["inicio"], format="DD/MM/YYYY")
        if (meta, inicio) != (config["meta"], config["inicio"]):
            if st.button("Salvar metas", width="stretch"):
                if salvar(config={"meta": meta, "inicio": inicio}):
                    st.rerun()
        st.caption(f"💾 Último salvamento: {st.session_state.get('ultimo_salvo', '—')}")
        if st.button("🔄 Recarregar planilha", width="stretch",
                     help="Use depois de editar a planilha direto no Excel."):
            st.cache_data.clear()
            st.session_state.versao = st.session_state.get("versao", 0) + 1
            st.rerun()
        st.caption("Feche a planilha no Excel enquanto usa o app, senão o salvamento é bloqueado.")


# ═════════════ Página: Início ═════════════
def pagina_inicio():
    st.markdown('<div class="titulo">Bom estudo, <span>Chris</span></div>', unsafe_allow_html=True)
    chips = "".join(f'<span class="chip">{c["icone"]} {k} · <b>{(c["prova"] - hoje).days} dias</b> para a prova</span>'
                    for k, c in CONCURSOS.items())
    st.markdown(chips, unsafe_allow_html=True)
    st.write("")

    horas = registros["Minutos"].sum() / 60 if len(registros) else 0.0
    planejadas = meta / 7 * max(1, (hoje - inicio).days + 1)
    feitos = topicos["Status"].isin(FEITO).sum()
    pct = feitos / len(topicos) * 100 if len(topicos) else 0
    acerto = taxa_acerto(topicos)
    k1, k2, k3 = st.columns(3)
    with k1:
        kpi("Horas estudadas × planejadas", fmt_h(horas), f"de {fmt_h(planejadas)} até hoje",
            horas / planejadas * 100 if planejadas else 0)
    with k2:
        kpi("Edital concluído", fmt_p(pct), f"{feitos} de {len(topicos)} tópicos", pct)
    with k3:
        kpi("Taxa de acertos", fmt_p(acerto), f"{int(topicos['Nº Questões'].sum())} questões", acerto)
    st.write("")

    esq, dir_ = st.columns([1, 1.15], gap="large")
    with esq:
        st.markdown("### 🍅 Pomodoro")
        pomodoro()
        hoje_reg = registros[registros["Data"] == hoje] if len(registros) else registros
        mins_hoje = hoje_reg["Minutos"].sum() if len(hoje_reg) else 0
        n_pomo = int((hoje_reg["Tipo"] == "Pomodoro").sum()) if len(hoje_reg) else 0
        st.caption(f"Hoje: **{fmt_h(mins_hoje / 60)}** de {fmt_h(meta / 7)} · {n_pomo} pomodoro(s)")
        registrar_estudo(topicos, "inicio")
    with dir_:
        st.markdown("### 🎯 Plano de hoje")
        dias = ["Segunda-feira", "Terça-feira", "Quarta-feira", "Quinta-feira", "Sexta-feira", "Sábado", "Domingo"]
        dia = dias[hoje.weekday()]
        if hoje <= CONCURSOS["PMPE"]["prova"] and dia in grade:
            linha = grade[dia]
            itens = [(k, str(v)) for k, v in linha.items()
                     if k and k not in ("Dia da Semana", "Status Semanal", "Observações / Ajustes")
                     and not vazio(v) and str(v).strip() != "-"]
            feitos_plano = sum(plano_feito.get((hoje, f"{k}: {v}"), False) for k, v in itens)
            st.caption(f"{dia} · Fase 1 · {feitos_plano} de {len(itens)} feitos")
            for i, (k, v) in enumerate(itens):
                item = f"{k}: {v}"
                chave = f"plano_{hoje}_{i}_v{st.session_state.get('versao', 0)}"
                st.checkbox(f"{v} · _{k}_",
                            value=plano_feito.get((hoje, item), False), key=chave,
                            on_change=lambda it=item, ch=chave: salvar(plano=(hoje, it, st.session_state[ch])))
            obs = linha.get("Observações / Ajustes")
            if not vazio(obs):
                st.caption(f"💡 {obs}")
        elif hoje <= CONCURSOS["PCPE"]["prova"]:
            st.info("Fase 2 (22/02 a 06/03): imersão PCPE no estilo Cebraspe (Certo/Errado).")
        else:
            st.info("Provas concluídas. Mantenha as revisões em dia.")
        if n_rev:
            st.button(f"🔁 Você tem {n_rev} revisão(ões) para hoje", on_click=ir, args=("revisoes",), width="stretch")

    st.write("")
    with st.expander("📊 Ver gráficos de desempenho"):
        res = resumo_materias(topicos)
        mapa = cores(res)
        g1, g2 = st.columns(2)
        with g1:
            fig = px.bar(res, x="% Concluído", y="Matéria", orientation="h", color="Matéria",
                         color_discrete_map=mapa, text="% Concluído", range_x=[0, 100])
            fig.update_traces(texttemplate="%{text:.0f}%", textposition="outside")
            fig.update_yaxes(autorange="reversed", title=None)
            st.plotly_chart(estilo(fig, titulo="Progresso do edital por matéria"), width="stretch")
        with g2:
            com_q = res[res["questoes"] > 0]
            if len(com_q):
                fig = px.bar(com_q, x="Matéria", y="% Acerto", color="Matéria", color_discrete_map=mapa,
                             text="% Acerto", range_y=[0, 105])
                fig.add_hline(y=70, line_dash="dot", line_color=VINHO_CLARO, annotation_text="meta 70%",
                              annotation_font_color=VINHO_CLARO)
                fig.update_traces(texttemplate="%{text:.0f}%", textposition="outside")
                fig.update_xaxes(title=None, tickangle=-25)
                st.plotly_chart(estilo(fig, titulo="Taxa de acertos por matéria"), width="stretch")
            else:
                st.caption("Preencha **Questões** e **% Acerto** nas tabelas para ver a taxa de acertos por matéria.")
        g3, g4 = st.columns(2)
        with g3:
            if len(registros):
                nome_cor = {MATERIAS[a][0]: MATERIAS[a][1] for a in MATERIAS}
                hm = registros.groupby("Matéria", as_index=False)["Minutos"].sum()
                hm["Horas"] = (hm["Minutos"] / 60).round(1)
                fig = px.bar(hm.sort_values("Horas"), x="Horas", y="Matéria", orientation="h", color="Matéria",
                             color_discrete_map=nome_cor, text="Horas")
                fig.update_yaxes(title=None)
                st.plotly_chart(estilo(fig, titulo="Horas estudadas por matéria"), width="stretch")
            else:
                st.caption("Registre pomodoros para ver as horas por matéria.")
        with g4:
            dias14 = pd.date_range(hoje - dt.timedelta(days=13), hoje).date
            por_dia = registros.groupby("Data")["Minutos"].sum() / 60 if len(registros) else pd.Series(dtype=float)
            serie = pd.DataFrame({"Dia": dias14, "Horas": [round(por_dia.get(d, 0.0), 2) for d in dias14]})
            fig = go.Figure(go.Bar(x=serie["Dia"], y=serie["Horas"], marker_color=VINHO_CLARO))
            fig.add_hline(y=meta / 7, line_dash="dot", line_color=APAGADO,
                          annotation_text=f"meta diária {meta / 7:.1f}h", annotation_font_color=APAGADO)
            fig.update_xaxes(tickformat="%d/%m")
            st.plotly_chart(estilo(fig, titulo="Horas por dia · últimos 14 dias"), width="stretch")


# ═════════════ Página: concurso (PMPE ou PCPE) ═════════════
def pagina_concurso(chave):
    conc = CONCURSOS[chave]
    df = topicos[topicos[chave]]
    mat = st.session_state.get("materia_sel")
    st.markdown(f'<div class="titulo">{conc["icone"]} <span>{conc["nome"]}</span></div>', unsafe_allow_html=True)
    st.markdown(f"<span class='chip'>Banca <b>{conc['banca']}</b></span>"
                f"<span class='chip'>Prova <b>{conc['prova']:%d/%m/%Y}</b> · faltam {(conc['prova'] - hoje).days} dias</span>"
                f"<span class='chip'>Tópicos comuns aos dois concursos atualizam nos dois</span>", unsafe_allow_html=True)
    st.write("")
    f_ = df["Status"].isin(FEITO).sum()
    p_ = f_ / len(df) * 100
    a1, a2, a3 = st.columns(3)
    with a1:
        kpi("Edital concluído", fmt_p(p_), f"{f_} de {len(df)} tópicos", p_)
    with a2:
        kpi("Taxa de acertos", fmt_p(taxa_acerto(df)), f"{int(df['Nº Questões'].sum())} questões", taxa_acerto(df))
    with a3:
        kpi("Revisões para hoje", f"{len(revisoes_pendentes(df, hoje))}", "deste concurso")
    st.write("")

    if mat:
        vis = df[df["Matéria"] == mat]
        st.markdown(f"### {mat}")
    else:
        vis = df
        res_c = resumo_materias(df)
        fig = px.bar(res_c, x="Matéria", y="% Concluído", color="Matéria", color_discrete_map=cores(res_c),
                     text="% Concluído", range_y=[0, 110])
        fig.update_traces(texttemplate="%{text:.0f}%", textposition="outside")
        fig.update_xaxes(title=None, tickangle=-20)
        st.plotly_chart(estilo(fig, 260, "Progresso por matéria"), width="stretch", key=f"graf_{chave}")

    filtro = st.pills("Status", STATUS, selection_mode="multi", default=STATUS, key=f"st_{chave}_{mat}")
    vis = vis[vis["Status"].isin(filtro or STATUS)]
    st.caption("✏️ Clique numa célula para editar. Cada alteração é salva sozinha na planilha original.")
    if len(vis):
        tabela_editavel(vis, f"tab_{chave}_{mat}")
        c1, c2 = st.columns([3, 1])
        abrir = c1.selectbox("Estudar um tópico desta lista", list(vis.index),
                             format_func=lambda i: f"{vis.at[i, 'Cód']} · {vis.at[i, 'Tópico']}",
                             key=f"abrir_{chave}_{mat}", label_visibility="collapsed")
        c2.button("📚 Estudar este tópico", width="stretch", on_click=ir, args=("estudar",),
                  kwargs={"estudo_id": abrir}, key=f"btn_abrir_{chave}_{mat}")
    else:
        st.caption("Nenhum tópico com esses filtros.")


# ═════════════ Materiais do curso (videoaulas e PDFs) ═════════════
def materiais(cod):
    nota = ler_anotacoes().get(cod, {})
    links = nota.get("links", [])
    st.markdown("#### 🎬 Videoaulas e links")
    if links:
        for n, link in enumerate(links, 1):
            st.link_button(f"▶ Abrir aula {n} · {link[:60]}", link, width="stretch")
    else:
        st.caption("Nenhum link ainda. Copie o endereço da aula no Gran Cursos e cole abaixo.")
    with st.form(f"link_{cod}", clear_on_submit=True):
        novo = st.text_input("Colar link da videoaula", placeholder="https://...")
        if st.form_submit_button("➕ Adicionar link") and novo.strip():
            salvar_anotacao(cod, links=links + [novo.strip()])
            st.rerun()

    st.markdown("#### 📄 PDFs")
    PASTA_MATERIAIS.mkdir(exist_ok=True)
    todos_pdfs = sorted(p.relative_to(PASTA_MATERIAIS).as_posix() for p in PASTA_MATERIAIS.rglob("*.pdf"))
    # Entram sozinhos os PDFs com o código do tópico no nome (ex.: "PENAL-08 Furto e Roubo.pdf").
    marcador = cod.replace("/", "-").lower()
    vinculados = [p for p in todos_pdfs if marcador in p.lower()] + \
        [p for p in nota.get("pdfs", []) if p in todos_pdfs and marcador not in p.lower()]
    if not todos_pdfs:
        st.caption(f"Coloque os PDFs baixados do curso na pasta **materiais**, dentro de {BASE.name}. "
                   f"Se o nome do arquivo tiver o código {cod}, ele aparece aqui sozinho.")
    if vinculados:
        escolhido = st.selectbox("PDF", vinculados, key=f"pdf_sel_{cod}") if len(vinculados) > 1 else vinculados[0]
        caminho = PASTA_MATERIAIS / escolhido
        dados = caminho.read_bytes()
        try:
            st.pdf(dados, height=700)
        except Exception:
            st.caption("Para ver o PDF aqui dentro, rode uma vez: pip install streamlit-pdf")
        st.download_button(f"⬇️ Abrir {caminho.name}", dados, file_name=caminho.name, mime="application/pdf",
                           key=f"pdf_dl_{cod}")
    outros = [p for p in todos_pdfs if p not in vinculados]
    if outros:
        with st.expander("Vincular outro PDF da pasta materiais a este tópico"):
            extra = st.selectbox("Arquivo", outros, key=f"pdf_add_{cod}")
            if st.button("🔗 Vincular", key=f"pdf_btn_{cod}"):
                salvar_anotacao(cod, pdfs=nota.get("pdfs", []) + [extra])
                st.rerun()


# ═════════════ Página: estudar ═════════════
def pagina_estudar():
    st.markdown('<div class="titulo">📚 <span>Estudar</span></div>', unsafe_allow_html=True)
    ids = list(topicos.index)
    atual = st.session_state.get("estudo_id")
    if atual not in ids:
        atual = ids[0]
    c1, c2 = st.columns([1, 2])
    mats = list(dict.fromkeys(topicos["Matéria"]))
    mat = c1.selectbox("Matéria", mats, index=mats.index(topicos.at[atual, "Matéria"]), key=f"est_mat_{atual}")
    sub = topicos[topicos["Matéria"] == mat]
    if atual not in sub.index:
        atual = sub.index[0]
    escolhido = c2.selectbox("Tópico", list(sub.index), index=list(sub.index).index(atual),
                             format_func=lambda i: f"{sub.at[i, 'Cód']} · {sub.at[i, 'Tópico']}", key=f"est_top_{mat}_{atual}")
    if escolhido != st.session_state.get("estudo_id"):
        st.session_state.estudo_id = escolhido
        st.rerun()
    t = topicos.loc[escolhido]
    cod = t["Cód"]
    concursos = " · ".join(k for k in CONCURSOS if t[k])
    st.markdown(f"<span class='chip' style='border-color:{t['Cor']}'>{t['Matéria']}</span>"
                f"<span class='chip'>Cód <b>{cod}</b></span><span class='chip'>Cai em <b>{concursos}</b></span>"
                f"<span class='chip'>Status <b>{t['Status']}</b></span>", unsafe_allow_html=True)
    st.markdown(f"## {t['Tópico']}")

    conteudo = base_estudos.get(cod, {"resumo": "", "flashcards": []})
    analise = banca.get(cod, {"palavras": [], "AOCP": "", "Cebraspe": ""})
    aba_res, aba_pts, aba_banca, aba_mat, aba_fc, aba_nota, aba_prog = st.tabs(
        ["📖 Resumo", "📌 Pontos para anotar", "🎯 Como a banca cobra", "🎬 Materiais", "🃏 Flashcards",
         "✍️ Minhas anotações", "✅ Progresso"])
    with aba_mat:
        materiais(cod)
    with aba_res:
        if analise["palavras"]:
            st.markdown("🔑 " + "".join(f"<span class='chip'>{p}</span>" for p in analise["palavras"]),
                        unsafe_allow_html=True)
            st.write("")
        if conteudo["resumo"]:
            st.markdown(conteudo["resumo"])
            st.caption("Resumo de apoio. Confira sempre na lei seca e no edital.")
        else:
            st.info(f"Ainda não há resumo para {cod}. Escreva o seu em **Minhas anotações** ou acrescente em "
                    f"base_estudos/{t['aba']}.md.")
    with aba_pts:
        pontos = pontos_do_resumo(conteudo["resumo"])
        if pontos:
            st.caption("Os pontos que mais caem neste tópico. Copie para as suas anotações e complete com as suas palavras.")
            for n, ponto in enumerate(pontos, 1):
                st.markdown(f"**{n}.** {ponto}")
            if st.button("➕ Copiar pontos para Minhas anotações", key=f"copiar_{cod}"):
                nota_atual = ler_anotacoes().get(cod, {})
                texto_novo = (nota_atual.get("texto", "").rstrip() + "\n\n" if nota_atual.get("texto") else "") + \
                    "\n".join(f"- {p}" for p in pontos)
                salvar_anotacao(cod, texto=texto_novo)
                st.toast("Pontos copiados para Minhas anotações ✓", icon="📌")
        else:
            st.info("Ainda não há pontos para este tópico.")
    with aba_banca:
        for nome_banca, conc in (("AOCP", "PMPE"), ("Cebraspe", "PCPE")):
            if not t[conc]:
                continue
            st.markdown(f"#### {CONCURSOS[conc]['icone']} {nome_banca} · {conc}")
            st.markdown(f"<div class='resumo'><b>Neste assunto:</b> {analise[nome_banca] or 'sem análise cadastrada.'}"
                        f"<br><br><small style='color:{APAGADO}'>Estilo da banca: {PERFIL_BANCA[nome_banca]}</small></div>",
                        unsafe_allow_html=True)
            st.write("")
        st.caption("Tendências observadas em provas anteriores, e não garantia de cobrança. Resolva questões recentes da banca.")
    with aba_fc:
        cartas = conteudo["flashcards"]
        if not cartas:
            st.info("Ainda não há flashcards para este tópico.")
        else:
            k = f"fc_{cod}"
            st.session_state.setdefault(k, 0)
            i = st.session_state[k] % len(cartas)
            mostrar = st.session_state.get(f"{k}_ver", False)
            pergunta, resposta = cartas[i]
            st.markdown(f"<div class='flash'><small>Cartão {i + 1} de {len(cartas)}</small>{pergunta}</div>",
                        unsafe_allow_html=True)
            if mostrar:
                st.markdown(f"<div class='flash' style='margin-top:10px;border-color:{VINHO_CLARO}'>"
                            f"<small>Resposta</small>{resposta}</div>", unsafe_allow_html=True)
            st.write("")
            b1, b2, b3 = st.columns(3)
            b1.button("👁️ Mostrar resposta", width="stretch", key=f"{k}_b1",
                      on_click=lambda: st.session_state.update({f"{k}_ver": True}))
            b2.button("◀ Anterior", width="stretch", key=f"{k}_b2",
                      on_click=lambda: st.session_state.update({k: i - 1, f"{k}_ver": False}))
            b3.button("Próximo ▶", width="stretch", key=f"{k}_b3",
                      on_click=lambda: st.session_state.update({k: i + 1, f"{k}_ver": False}))
    with aba_nota:
        nota = ler_anotacoes().get(cod, {})
        with st.form(f"nota_{cod}"):
            texto = st.text_area("Suas anotações (aceita **negrito**, listas com - etc.)", nota.get("texto", ""),
                                 height=260)
            links = st.text_area("Links de aulas, PDFs e leis (um por linha)", "\n".join(nota.get("links", [])), height=90)
            if st.form_submit_button("💾 Salvar anotações"):
                salvar_anotacao(cod, texto=texto, links=[l.strip() for l in links.splitlines() if l.strip()])
                st.toast("Anotações salvas ✓", icon="✍️")
        if nota.get("links"):
            st.markdown("**Links salvos:**  \n" + "  \n".join(f"- {l}" for l in nota["links"]))
    with aba_prog:
        n1, n2 = st.columns(2)
        novo_status = n1.selectbox("Status do tópico", STATUS, index=STATUS.index(t["Status"]), key=f"status_{cod}")
        if novo_status != t["Status"]:
            alt = [(escolhido, "Status", novo_status)]
            if novo_status in FEITO and not t["Data Estudada"]:
                alt.append((escolhido, "Data Estudada", hoje))
            if salvar(alt):
                avisar(f"{cod} marcado como {novo_status} ✓", "✅")
                st.rerun()
        n2.metric("Questões · % acerto", f"{t['Nº Questões']} · {t['% Acerto']:.0f}%")
        registrar_estudo(topicos, f"est_{cod}", cod_fixo=cod)

    pos = ids.index(escolhido)
    st.write("")
    p1, _, p2 = st.columns([1, 2, 1])
    if pos > 0:
        p1.button("◀ Tópico anterior", width="stretch", on_click=ir, args=("estudar",), kwargs={"estudo_id": ids[pos - 1]})
    if pos < len(ids) - 1:
        p2.button("Próximo tópico ▶", width="stretch", on_click=ir, args=("estudar",), kwargs={"estudo_id": ids[pos + 1]})


# ═════════════ Página: revisões ═════════════
def pagina_revisoes():
    st.markdown('<div class="titulo">🔁 <span>Revisões</span></div>', unsafe_allow_html=True)
    st.caption("As revisões de 24h, 7, 15 e 30 dias contam a partir de **Estudado em**. "
               "Marque ✔ quando fizer; fica gravado como 'Feito' na planilha.")
    pend = revisoes_pendentes(topicos, hoje)
    if pend.empty:
        st.success("Nenhuma revisão vencida. 👏")
    else:
        pend = pend.sort_values("Vence em").reset_index(drop=True)
        pend.insert(0, "Feito", False)
        ed = st.data_editor(
            pend, hide_index=True, width="stretch", key=f"rev_v{st.session_state.get('versao', 0)}",
            column_order=["Feito", "Matéria", "Cód", "Tópico", "Revisão", "Vence em", "Atraso (dias)"],
            disabled=["Matéria", "Cód", "Tópico", "Revisão", "Vence em", "Atraso (dias)"],
            column_config={"Feito": st.column_config.CheckboxColumn("✔", width="small"),
                           "Vence em": st.column_config.DateColumn("Vence em", format="DD/MM/YYYY"),
                           "Tópico": st.column_config.TextColumn(width="large")})
        marcadas = ed[ed["Feito"]]
        if len(marcadas):
            alt = [(r["id"], r["Coluna"], "Feito") for _, r in marcadas.iterrows()]
            for ident in marcadas["id"].unique():
                feitas_agora = set(marcadas.loc[marcadas["id"] == ident, "Coluna"])
                faltam = [c for c, _ in REVISOES if topicos.at[ident, c] == "Pendente" and c not in feitas_agora]
                if not faltam and topicos.at[ident, "Status"] == "Concluído":
                    alt.append((ident, "Status", "Revisado"))
            if salvar(alt):
                avisar("Revisão registrada ✓", "🔁")
                st.rerun()
    prox = revisoes_pendentes(topicos, hoje + dt.timedelta(days=7))
    prox = prox[prox["Vence em"] > hoje] if not prox.empty else prox
    if not prox.empty:
        with st.expander(f"Próximos 7 dias · {len(prox)} revisões"):
            st.dataframe(prox.sort_values("Vence em")[["Vence em", "Matéria", "Cód", "Tópico", "Revisão"]],
                         hide_index=True, width="stretch",
                         column_config={"Vence em": st.column_config.DateColumn(format="DD/MM/YYYY")})


# ═════════════ Página: resumo do dia ═════════════
def topicos_do_dia():
    """Tópicos planejados para hoje, estudados hoje ou com estudo registrado hoje."""
    ids = set(topicos.index[(topicos["Data Planejada"] == hoje) | (topicos["Data Estudada"] == hoje)])
    if len(registros):
        cods = set(registros.loc[registros["Data"] == hoje, "Cód"].dropna().astype(str))
        ids |= set(topicos.index[topicos["Cód"].isin(cods)])
    return [i for i in topicos.index if i in ids]


def pagina_resumo_dia():
    st.markdown('<div class="titulo">📝 Resumo <span>do dia</span></div>', unsafe_allow_html=True)
    st.caption(f"{hoje:%d/%m/%Y} · Junta resumo, pontos para anotar, palavras-chave e como a banca cobra, separados por matéria.")
    sugeridos = topicos_do_dia()
    escolha = st.multiselect(
        "Assuntos de hoje", list(topicos.index), default=sugeridos, key=f"resumo_dia_{hoje}",
        format_func=lambda i: f"{topicos.at[i, 'Matéria'].split(' ', 1)[0]} {topicos.at[i, 'Cód']} · {topicos.at[i, 'Tópico'][:60]}",
        placeholder="Escolha os assuntos que você estudou ou vai estudar hoje")
    if not sugeridos:
        st.caption("Dica: os assuntos entram aqui sozinhos quando têm data planejada ou estudada hoje, "
                   "ou quando você registra um pomodoro neles.")
    if not escolha:
        return
    texto_md = [f"# Resumo do dia · {hoje:%d/%m/%Y}\n"]
    sel = topicos.loc[escolha]
    for mat, grupo in sel.groupby("Matéria", sort=False):
        st.markdown(f"## {mat}")
        texto_md.append(f"\n## {mat}\n")
        for ident, t in grupo.iterrows():
            cod = t["Cód"]
            conteudo = base_estudos.get(cod, {"resumo": ""})
            analise = banca.get(cod, {"palavras": [], "AOCP": "", "Cebraspe": ""})
            with st.expander(f"{cod} · {t['Tópico']}", expanded=True):
                if analise["palavras"]:
                    st.markdown("🔑 " + "".join(f"<span class='chip'>{p}</span>" for p in analise["palavras"]),
                                unsafe_allow_html=True)
                pontos = pontos_do_resumo(conteudo["resumo"])
                if pontos:
                    st.markdown("**📌 Pontos para anotar**\n" + "\n".join(f"- {p}" for p in pontos))
                for nome_banca, conc in (("AOCP", "PMPE"), ("Cebraspe", "PCPE")):
                    if t[conc] and analise[nome_banca]:
                        st.markdown(f"**🎯 {nome_banca}:** {analise[nome_banca]}")
                st.button("📚 Abrir para estudar", key=f"rd_abrir_{ident}", on_click=ir, args=("estudar",),
                          kwargs={"estudo_id": ident})
            texto_md.append(f"\n### {cod} · {t['Tópico']}\n")
            if analise["palavras"]:
                texto_md.append("Palavras-chave: " + "; ".join(analise["palavras"]) + "\n")
            texto_md += [f"- {p}" for p in pontos]
            for nome_banca, conc in (("AOCP", "PMPE"), ("Cebraspe", "PCPE")):
                if t[conc] and analise[nome_banca]:
                    texto_md.append(f"\nComo a {nome_banca} cobra: {analise[nome_banca]}")
    st.download_button("⬇️ Baixar resumo do dia", "\n".join(texto_md), file_name=f"resumo_{hoje:%Y-%m-%d}.md",
                       mime="text/markdown")


if pagina.startswith("conc:"):
    pagina_concurso(pagina.split(":", 1)[1])
elif pagina == "estudar":
    pagina_estudar()
elif pagina == "resumo_dia":
    pagina_resumo_dia()
elif pagina == "revisoes":
    pagina_revisoes()
else:
    pagina_inicio()
